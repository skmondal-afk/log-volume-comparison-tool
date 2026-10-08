"""
components.py

Reads the list of components to check from an .xlsx file
(config.COMPONENTS_XLSX_PATH). Expects a header row with a "component"
column (case-insensitive); falls back to the first column if no such
header is found.

Requires the "openpyxl" package (pip install openpyxl).
"""

from openpyxl import load_workbook


def load_components(xlsx_path: str) -> list:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = wb.active

    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []

    col_index = 0
    for i, cell in enumerate(header):
        if cell and str(cell).strip().lower() == "component":
            col_index = i
            break

    components = []
    for row in rows:
        if len(row) > col_index and row[col_index] not in (None, ""):
            components.append(str(row[col_index]).strip())

    return components
